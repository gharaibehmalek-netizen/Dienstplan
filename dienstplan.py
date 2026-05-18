#!/usr/bin/env python3
"""
Dienstplan-Generator fuer Zahnarztpraxen
Verwaltet Mitarbeiter, Urlaub, Krankheit und erstellt Monatsplaene.
"""

import json
import os
import sys
import calendar
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

DATA_FILE = Path("praxis_daten.json")

SCHICHTEN = {
    "V": {"name": "Vormittag",  "zeiten": "08:00–14:00"},
    "N": {"name": "Nachmittag", "zeiten": "14:00–20:00"},
    "G": {"name": "Ganztag",    "zeiten": "08:00–20:00"},
    "U": {"name": "Urlaub",     "zeiten": ""},
    "K": {"name": "Krankheit",  "zeiten": ""},
    "F": {"name": "Frei",       "zeiten": ""},
}

FARBEN = {
    "V": "C6EFCE",  # hellgruen
    "N": "BDD7EE",  # hellblau
    "G": "FFEB9C",  # hellgelb
    "U": "F4CCCC",  # hellrot
    "K": "FCE4D6",  # lachsfarben
    "F": "EFEFEF",  # grau
    "WE": "D9D9D9", # wochenende
}

WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


def lade_daten() -> dict:
    if DATA_FILE.exists():
        with open(DATA_FILE, encoding="utf-8") as f:
            return json.load(f)
    return {"mitarbeiter": [], "abwesenheiten": {}, "plaene": {}}


def speichere_daten(daten: dict):
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(daten, f, ensure_ascii=False, indent=2)


def mitarbeiter_liste(daten: dict):
    ma = daten.get("mitarbeiter", [])
    if not ma:
        print("  Keine Mitarbeiter erfasst.")
        return
    print(f"\n  {'Nr':>3}  {'Name':<20}  {'Rolle':<20}  {'Std/Woche':>9}")
    print("  " + "-" * 57)
    for i, m in enumerate(ma, 1):
        print(f"  {i:>3}  {m['name']:<20}  {m['rolle']:<20}  {m.get('stunden_woche', 40):>9}")


def mitarbeiter_hinzufuegen(daten: dict):
    print("\n--- Neuen Mitarbeiter anlegen ---")
    name = input("  Name: ").strip()
    if not name:
        print("  Abgebrochen.")
        return
    rolle = input("  Rolle (z.B. Zahnärztin, ZFA, Empfang): ").strip() or "ZFA"
    try:
        stunden = int(input("  Wochenstunden [40]: ").strip() or "40")
    except ValueError:
        stunden = 40
    daten["mitarbeiter"].append({"name": name, "rolle": rolle, "stunden_woche": stunden})
    speichere_daten(daten)
    print(f"  ✓ {name} wurde hinzugefügt.")


def mitarbeiter_loeschen(daten: dict):
    mitarbeiter_liste(daten)
    try:
        nr = int(input("\n  Nummer löschen (0=Abbrechen): "))
    except ValueError:
        return
    if nr == 0:
        return
    ma = daten["mitarbeiter"]
    if 1 <= nr <= len(ma):
        entfernt = ma.pop(nr - 1)
        speichere_daten(daten)
        print(f"  ✓ {entfernt['name']} entfernt.")
    else:
        print("  Ungültige Nummer.")


def abwesenheit_erfassen(daten: dict, art: str):
    """art = 'U' (Urlaub) oder 'K' (Krankheit)"""
    mitarbeiter_liste(daten)
    ma = daten["mitarbeiter"]
    if not ma:
        return
    try:
        nr = int(input(f"\n  Mitarbeiter-Nummer: "))
    except ValueError:
        return
    if not (1 <= nr <= len(ma)):
        print("  Ungültige Nummer.")
        return
    person = ma[nr - 1]["name"]

    print("  Datum eingeben (JJJJ-MM-TT), Bereich mit '-' trennen z.B. 2025-06-02-2025-06-06")
    eintrag = input("  Datum: ").strip()

    try:
        if "-" in eintrag[10:]:  # Bereich
            teile = eintrag.split("-")
            von = date(int(teile[0]), int(teile[1]), int(teile[2]))
            bis = date(int(teile[3]), int(teile[4]), int(teile[5]))
        else:
            von = bis = date.fromisoformat(eintrag)
    except Exception:
        print("  Ungültiges Datumsformat.")
        return

    abw = daten.setdefault("abwesenheiten", {})
    tage_eingetragen = 0
    tag = von
    while tag <= bis:
        if tag.weekday() < 5:  # nur Mo-Fr
            key = tag.isoformat()
            if key not in abw:
                abw[key] = {}
            abw[key][person] = art
            tage_eingetragen += 1
        tag += timedelta(days=1)

    speichere_daten(daten)
    art_name = "Urlaub" if art == "U" else "Krankheit"
    print(f"  ✓ {art_name} für {person}: {von} bis {bis} ({tage_eingetragen} Tag(e)) erfasst.")


def abwesenheit_loeschen(daten: dict):
    print("  Datum (JJJJ-MM-TT): ", end="")
    eintrag = input().strip()
    mitarbeiter_liste(daten)
    ma = daten["mitarbeiter"]
    try:
        nr = int(input("  Mitarbeiter-Nummer: "))
        person = ma[nr - 1]["name"]
    except (ValueError, IndexError):
        return
    abw = daten.get("abwesenheiten", {})
    if eintrag in abw and person in abw[eintrag]:
        del abw[eintrag][person]
        if not abw[eintrag]:
            del abw[eintrag]
        speichere_daten(daten)
        print(f"  ✓ Eintrag für {person} am {eintrag} gelöscht.")
    else:
        print("  Kein Eintrag gefunden.")


def abwesenheiten_anzeigen(daten: dict):
    abw = daten.get("abwesenheiten", {})
    if not abw:
        print("  Keine Abwesenheiten erfasst.")
        return
    try:
        monat_str = input("  Monat anzeigen (JJJJ-MM, leer = alle): ").strip()
    except EOFError:
        monat_str = ""
    print()
    for tag in sorted(abw.keys()):
        if monat_str and not tag.startswith(monat_str):
            continue
        d = date.fromisoformat(tag)
        wt = WOCHENTAGE[d.weekday()]
        for person, art in abw[tag].items():
            art_name = SCHICHTEN.get(art, {}).get("name", art)
            print(f"  {wt} {tag}  {person:<20}  {art_name}")


def erstelle_monatsplan(daten: dict):
    print("\n--- Monatsplan erstellen ---")
    try:
        jahr = int(input("  Jahr [2025]: ").strip() or "2025")
        monat = int(input("  Monat (1-12): ").strip())
    except ValueError:
        print("  Ungültige Eingabe.")
        return

    ma_liste = daten.get("mitarbeiter", [])
    if not ma_liste:
        print("  Keine Mitarbeiter erfasst. Bitte zuerst Mitarbeiter anlegen.")
        return

    abw = daten.get("abwesenheiten", {})
    _, tage_im_monat = calendar.monthrange(jahr, monat)

    plaene = daten.setdefault("plaene", {})
    plan_key = f"{jahr}-{monat:02d}"
    plan = plaene.setdefault(plan_key, {})

    # Standard-Schichten fuer jeden Mitarbeiter pro Tag setzen
    for tag_nr in range(1, tage_im_monat + 1):
        tag = date(jahr, monat, tag_nr)
        tag_str = tag.isoformat()
        ist_wochenende = tag.weekday() >= 5

        for ma in ma_liste:
            person = ma["name"]
            if tag_str not in plan:
                plan[tag_str] = {}
            if person not in plan[tag_str]:
                # Abwesenheit pruefen
                abw_art = abw.get(tag_str, {}).get(person)
                if abw_art:
                    plan[tag_str][person] = abw_art
                elif ist_wochenende:
                    plan[tag_str][person] = "F"
                else:
                    plan[tag_str][person] = "G"  # Standard: Ganztag

    speichere_daten(daten)
    print(f"  ✓ Plan für {monat:02d}/{jahr} initialisiert.")

    # Schichten anpassen?
    antwort = input("  Schichten jetzt bearbeiten? (j/n) [n]: ").strip().lower()
    if antwort == "j":
        schichten_bearbeiten(daten, jahr, monat)

    # Export
    antwort = input("  Plan exportieren? (excel/text/nein) [text]: ").strip().lower() or "text"
    if antwort == "excel":
        exportiere_excel(daten, jahr, monat)
    elif antwort == "text":
        drucke_monatsplan(daten, jahr, monat)


def schichten_bearbeiten(daten: dict, jahr: int, monat: int):
    plan_key = f"{jahr}-{monat:02d}"
    plan = daten.get("plaene", {}).get(plan_key, {})
    ma_liste = daten.get("mitarbeiter", [])

    print("\n  Schichtcodes: V=Vormittag(8-14)  N=Nachmittag(14-20)  G=Ganztag(8-20)  U=Urlaub  K=Krankheit  F=Frei")
    print("  Eingabe: JJJJ-MM-TT [leer=Ende]")

    while True:
        tag_str = input("\n  Datum (leer=Ende): ").strip()
        if not tag_str:
            break
        try:
            tag = date.fromisoformat(tag_str)
        except ValueError:
            print("  Ungültiges Datum.")
            continue
        if tag.year != jahr or tag.month != monat:
            print("  Datum liegt außerhalb des Monats.")
            continue

        mitarbeiter_liste(daten)
        try:
            nr = int(input("  Mitarbeiter-Nummer: "))
            person = ma_liste[nr - 1]["name"]
        except (ValueError, IndexError):
            continue

        print(f"  Aktuell: {plan.get(tag_str, {}).get(person, '?')}")
        code = input("  Neuer Code (V/N/G/U/K/F): ").strip().upper()
        if code not in SCHICHTEN:
            print("  Ungültiger Code.")
            continue
        if tag_str not in plan:
            plan[tag_str] = {}
        plan[tag_str][person] = code

    speichere_daten(daten)
    print("  ✓ Schichten gespeichert.")


def drucke_monatsplan(daten: dict, jahr: int, monat: int):
    plan_key = f"{jahr}-{monat:02d}"
    plan = daten.get("plaene", {}).get(plan_key)
    ma_liste = daten.get("mitarbeiter", [])

    if not plan:
        print("  Kein Plan vorhanden. Bitte zuerst Plan erstellen.")
        return

    monat_name = ["", "Januar", "Februar", "März", "April", "Mai", "Juni",
                  "Juli", "August", "September", "Oktober", "November", "Dezember"][monat]
    print(f"\n{'='*80}")
    print(f"  DIENSTPLAN {monat_name} {jahr}  |  Zahnarztpraxis  |  Mo–Fr  08:00–20:00 Uhr")
    print(f"{'='*80}")

    namen = [m["name"] for m in ma_liste]
    breite = max(len(n) for n in namen) if namen else 10

    # Kopfzeile
    header = f"  {'Datum':<12} {'Tag':<4}"
    for n in namen:
        header += f" {n[:breite]:<{breite}}"
    print(header)
    print("  " + "-" * (16 + (breite + 1) * len(namen)))

    _, tage = calendar.monthrange(jahr, monat)
    for tag_nr in range(1, tage + 1):
        tag = date(jahr, monat, tag_nr)
        tag_str = tag.isoformat()
        wt = WOCHENTAGE[tag.weekday()]
        ist_we = tag.weekday() >= 5

        zeile = f"  {tag_str:<12} {wt:<4}"
        for person in namen:
            code = plan.get(tag_str, {}).get(person, "-")
            zeile += f" {code:<{breite}}"

        if ist_we:
            zeile += "  (Wochenende)"
        print(zeile)

    print(f"{'='*80}")
    print("\n  Legende: V=Vormittag(8-14)  N=Nachmittag(14-20)  G=Ganztag(8-20)  U=Urlaub  K=Krankheit  F=Frei")

    # Statistik
    print(f"\n  Statistik {monat_name} {jahr}:")
    print(f"  {'Name':<{breite+2}} {'G':>4} {'V':>4} {'N':>4} {'U':>4} {'K':>4} {'F':>4}")
    print("  " + "-" * (breite + 28))
    for person in namen:
        zaehler = {k: 0 for k in SCHICHTEN}
        for tag_str, belegung in plan.items():
            code = belegung.get(person, "F")
            zaehler[code] = zaehler.get(code, 0) + 1
        print(f"  {person:<{breite+2}} {zaehler['G']:>4} {zaehler['V']:>4} "
              f"{zaehler['N']:>4} {zaehler['U']:>4} {zaehler['K']:>4} {zaehler['F']:>4}")


def exportiere_excel(daten: dict, jahr: int, monat: int):
    if not EXCEL_AVAILABLE:
        print("  openpyxl nicht installiert. Bitte: pip install openpyxl")
        drucke_monatsplan(daten, jahr, monat)
        return

    plan_key = f"{jahr}-{monat:02d}"
    plan = daten.get("plaene", {}).get(plan_key)
    ma_liste = daten.get("mitarbeiter", [])

    if not plan:
        print("  Kein Plan vorhanden.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    monat_name = ["", "Januar", "Februar", "März", "April", "Mai", "Juni",
                  "Juli", "August", "September", "Oktober", "November", "Dezember"][monat]
    ws.title = f"{monat_name} {jahr}"

    # Titelzeile
    ws.merge_cells(f"A1:{get_column_letter(3 + len(ma_liste))}1")
    title_cell = ws["A1"]
    title_cell.value = f"Dienstplan {monat_name} {jahr} – Zahnarztpraxis  |  Mo–Fr  08:00–20:00 Uhr"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Kopfzeile
    headers = ["Datum", "Wochentag", "KW"] + [m["name"] for m in ma_liste]
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # Daten
    _, tage = calendar.monthrange(jahr, monat)
    for tag_nr in range(1, tage + 1):
        tag = date(jahr, monat, tag_nr)
        tag_str = tag.isoformat()
        row = tag_nr + 2
        ist_we = tag.weekday() >= 5
        kw = tag.isocalendar()[1]

        ws.cell(row=row, column=1, value=tag_str).border = border
        ws.cell(row=row, column=2, value=WOCHENTAGE[tag.weekday()]).border = border
        ws.cell(row=row, column=3, value=kw).border = border

        for col_off, ma in enumerate(ma_liste):
            person = ma["name"]
            code = plan.get(tag_str, {}).get(person, "F")
            col = col_off + 4
            cell = ws.cell(row=row, column=col)
            cell.value = code
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

            farbe = FARBEN.get("WE" if ist_we else code, "FFFFFF")
            cell.fill = PatternFill("solid", fgColor=farbe)

        # Zeilenhintergrund Wochenende
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = PatternFill(
                "solid", fgColor=FARBEN["WE"] if ist_we else "FFFFFF"
            )

    # Statistik-Blatt
    ws_stat = wb.create_sheet("Statistik")
    ws_stat["A1"] = f"Statistik {monat_name} {jahr}"
    ws_stat["A1"].font = Font(bold=True, size=12)

    stat_header = ["Name", "Rolle", "Ganztag", "Vormittag", "Nachmittag", "Urlaub", "Krankheit", "Frei", "Arbeitstage"]
    for col, h in enumerate(stat_header, 1):
        cell = ws_stat.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for row_off, ma in enumerate(ma_liste):
        person = ma["name"]
        zaehler = {"G": 0, "V": 0, "N": 0, "U": 0, "K": 0, "F": 0}
        for tag_str, belegung in plan.items():
            code = belegung.get(person, "F")
            if code in zaehler:
                zaehler[code] += 1
        arbeitstage = zaehler["G"] + zaehler["V"] + zaehler["N"]
        row = row_off + 3
        ws_stat.cell(row=row, column=1, value=person)
        ws_stat.cell(row=row, column=2, value=ma.get("rolle", ""))
        ws_stat.cell(row=row, column=3, value=zaehler["G"])
        ws_stat.cell(row=row, column=4, value=zaehler["V"])
        ws_stat.cell(row=row, column=5, value=zaehler["N"])
        ws_stat.cell(row=row, column=6, value=zaehler["U"])
        ws_stat.cell(row=row, column=7, value=zaehler["K"])
        ws_stat.cell(row=row, column=8, value=zaehler["F"])
        ws_stat.cell(row=row, column=9, value=arbeitstage)

    # Legende
    ws_leg = wb.create_sheet("Legende")
    ws_leg["A1"] = "Schichtcodes"
    ws_leg["A1"].font = Font(bold=True)
    for row, (code, info) in enumerate(SCHICHTEN.items(), 2):
        ws_leg.cell(row=row, column=1, value=code)
        ws_leg.cell(row=row, column=2, value=info["name"])
        ws_leg.cell(row=row, column=3, value=info["zeiten"])
        ws_leg.cell(row=row, column=1).fill = PatternFill("solid", fgColor=FARBEN.get(code, "FFFFFF"))

    # Spaltenbreiten anpassen
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 5
    for col in range(4, 4 + len(ma_liste)):
        ws.column_dimensions[get_column_letter(col)].width = 15

    dateiname = f"dienstplan_{plan_key}.xlsx"
    wb.save(dateiname)
    print(f"  ✓ Excel-Datei gespeichert: {dateiname}")


def zeige_plan(daten: dict):
    try:
        jahr = int(input("  Jahr: ").strip())
        monat = int(input("  Monat (1-12): ").strip())
    except ValueError:
        return
    drucke_monatsplan(daten, jahr, monat)
    antwort = input("\n  Als Excel exportieren? (j/n) [n]: ").strip().lower()
    if antwort == "j":
        exportiere_excel(daten, jahr, monat)


def hauptmenue():
    print("\n" + "=" * 50)
    print("  DIENSTPLAN – Zahnarztpraxis")
    print("=" * 50)
    print("  1  Mitarbeiter anzeigen")
    print("  2  Mitarbeiter hinzufügen")
    print("  3  Mitarbeiter löschen")
    print("  ─────────────────────────")
    print("  4  Urlaub erfassen")
    print("  5  Krankheit erfassen")
    print("  6  Abwesenheit löschen")
    print("  7  Abwesenheiten anzeigen")
    print("  ─────────────────────────")
    print("  8  Monatsplan erstellen/bearbeiten")
    print("  9  Monatsplan anzeigen / exportieren")
    print("  ─────────────────────────")
    print("  0  Beenden")
    print("=" * 50)


def main():
    daten = lade_daten()

    # Demo-Daten anlegen wenn noch leer
    if not daten["mitarbeiter"]:
        print("\n  Keine Daten vorhanden. Demo-Mitarbeiter werden angelegt...")
        daten["mitarbeiter"] = [
            {"name": "Dr. Schmidt",    "rolle": "Zahnärztin",  "stunden_woche": 40},
            {"name": "Dr. Müller",     "rolle": "Zahnarzt",    "stunden_woche": 40},
            {"name": "Anna Weber",     "rolle": "ZFA",         "stunden_woche": 40},
            {"name": "Lisa Braun",     "rolle": "ZFA",         "stunden_woche": 30},
            {"name": "Jana Koch",      "rolle": "Empfang",     "stunden_woche": 20},
        ]
        speichere_daten(daten)
        print("  ✓ 5 Demo-Mitarbeiter angelegt.")

    while True:
        hauptmenue()
        try:
            wahl = input("  Auswahl: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n  Auf Wiedersehen!")
            break

        if wahl == "0":
            print("  Auf Wiedersehen!")
            break
        elif wahl == "1":
            mitarbeiter_liste(daten)
        elif wahl == "2":
            mitarbeiter_hinzufuegen(daten)
        elif wahl == "3":
            mitarbeiter_loeschen(daten)
        elif wahl == "4":
            abwesenheit_erfassen(daten, "U")
        elif wahl == "5":
            abwesenheit_erfassen(daten, "K")
        elif wahl == "6":
            abwesenheit_loeschen(daten)
        elif wahl == "7":
            abwesenheiten_anzeigen(daten)
        elif wahl == "8":
            erstelle_monatsplan(daten)
        elif wahl == "9":
            zeige_plan(daten)
        else:
            print("  Ungültige Eingabe.")


if __name__ == "__main__":
    main()
